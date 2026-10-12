"""
core/templates.py — 서로이웃 신청 메시지 / 댓글 템플릿 관리
- 사용자가 직접 확인·수정한 내용만 사용하도록 저장/불러오기 지원
- 동일 메시지 대량 반복 사용을 방지하기 위한 랜덤 선택 헬퍼 포함
"""

import json
import os
import random
import re
import threading

TEMPLATES_FILE = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "templates.json"
)

_lock = threading.Lock()

DEFAULTS = {
    "neighbor_messages": [
        "안녕하세요.\n블로그를 둘러보다가 좋은 글이 많아서 서로 소통하면 좋을 것 같아 신청드립니다.\n앞으로 좋은 정보 서로 공유했으면 좋겠습니다.",
        "안녕하세요! 꾸준히 좋은 글 올려주시네요.\n서로이웃 맺고 소통해요 :)",
    ],
    "comments": [
        "정보 잘 보고 갑니다. 좋은 글 감사합니다.",
        "내용이 유익하네요. 잘 보고 갑니다.",
        "좋은 정보 감사합니다. 다음 글도 기대하겠습니다.",
    ],
}


def _load_all():
    if not os.path.exists(TEMPLATES_FILE):
        return dict(DEFAULTS)
    try:
        with open(TEMPLATES_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
        # 기본 키 보장
        for k, v in DEFAULTS.items():
            data.setdefault(k, v)
        return data
    except Exception:
        return dict(DEFAULTS)


def _save_all(data):
    try:
        with open(TEMPLATES_FILE, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        return True
    except Exception:
        return False


def get_neighbor_messages() -> list:
    return list(_load_all().get("neighbor_messages", []))


def get_comments() -> list:
    return list(_load_all().get("comments", []))


def save_neighbor_messages(messages: list) -> bool:
    with _lock:
        data = _load_all()
        data["neighbor_messages"] = [m for m in messages if m and m.strip()]
        return _save_all(data)


def save_comments(comments: list) -> bool:
    with _lock:
        data = _load_all()
        data["comments"] = [c for c in comments if c and c.strip()]
        return _save_all(data)


def add_neighbor_message(message: str) -> bool:
    msgs = get_neighbor_messages()
    if message and message.strip() and message not in msgs:
        msgs.append(message.strip())
        return save_neighbor_messages(msgs)
    return True


def add_comment(comment: str) -> bool:
    comments = get_comments()
    if comment and comment.strip() and comment not in comments:
        comments.append(comment.strip())
        return save_comments(comments)
    return True


def pick_random(kind: str, used: set) -> str:
    """
    사용하지 않은 템플릿 중 랜덤 선택.
    모두 사용했으면 used를 초기화하고 다시 선택.
    kind: "neighbor" | "comment"
    """
    if kind == "neighbor":
        pool = get_neighbor_messages()
    else:
        pool = get_comments()

    available = [t for t in pool if t not in used]
    if not available:
        used.clear()
        available = list(pool)
    if not available:
        return ""
    picked = random.choice(available)
    used.add(picked)  # 호출자의 used 집합에도 반영
    return picked


# ─── 랜덤 텍스트 생성 (입력창 내용 기반) ──────────────────
# 규칙:
#   1) 빈 줄(엔터 2번)로 구분된 여러 후보 중 하나가 랜덤 선택됨
#   2) {A|B|C} 치환 문법으로 문장 내 단어/문구가 랜덤 조합됨
#      예: "안녕하세요 {반갑습니다|잘 부탁드립니다}!"
_CANDIDATE_SEP_RE = re.compile(r"\n\s*\n+")
_OPTION_GROUP_RE = re.compile(r"\{([^{}]*)\}")


def split_candidates(text: str) -> list:
    """입력 텍스트를 빈 줄 기준으로 후보 리스트로 분리"""
    if not text:
        return []
    return [c.strip() for c in _CANDIDATE_SEP_RE.split(text.strip()) if c.strip()]


def expand_random_text(text: str) -> str:
    """{A|B|C} 치환 문법을 랜덤 조합으로 확장 (중첩 괄호도 처리)"""
    if not text or "{" not in text:
        return text

    prev = None
    while prev != text:
        prev = text
        text = _OPTION_GROUP_RE.sub(
            lambda m: random.choice([o.strip() for o in m.group(1).split("|")])
            if m.group(1).split("|") else "",
            text,
        )
    return text


def pick_from_text(text: str, used: set) -> str:
    """
    입력창에 직접 입력한 내용에서 랜덤 선택:
    - 빈 줄로 구분된 후보 중 사용하지 않은 것을 랜덤 선택 (모두 사용 시 초기화)
    - 선택된 후보에 {A|B} 치환 문법이 있으면 랜덤 조합으로 확장
    - used에는 확장 전 원문 후보가 기록됨 (중복 방지 기준)
    """
    candidates = split_candidates(text)
    if not candidates:
        return expand_random_text(text or "")

    available = [c for c in candidates if c not in used]
    if not available:
        used.clear()
        available = list(candidates)
    picked = random.choice(available)
    used.add(picked)
    return expand_random_text(picked)


# ─── 파일 첨부 불러오기 (메모장 / 엑셀) ──────────────────
TXT_EXTS = (".txt", ".md", ".log")
EXCEL_EXTS = (".xlsx", ".xlsm")


def read_template_file(path: str) -> str:
    """
    메모장(.txt/.md) 또는 엑셀(.xlsx) 파일에서 템플릿 내용을 읽어온다.
    - 텍스트: 파일 내용 그대로 (빈 줄 구분 후보 / {A|B} 치환 모두 적용 가능)
    - 엑셀: 첫 시트의 첫 열(또는 단일 행) 값을 하나씩 읽어 빈 줄로 연결
      → 각 행이 랜덤 후보 1개가 됨
    실패 시 FileNotFoundError 그대로 전파 (호출자에서 안내).
    """
    if not path or not os.path.exists(path):
        raise FileNotFoundError(path)

    lower = path.lower()
    if lower.endswith(EXCEL_EXTS):
        from openpyxl import load_workbook  # 지연 임포트 (엑셀 파일일 때만)

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            rows = []
            for row in ws.iter_rows(values_only=True):
                if not row:
                    continue
                # 첫 열 우선, 비어 있으면 해당 행의 다른 값 사용
                cell = next((v for v in row if v is not None and str(v).strip()), None)
                if cell is not None:
                    text = str(cell).strip()
                    if text:
                        rows.append(text)
            return "\n\n".join(rows)
        finally:
            wb.close()

    # 텍스트 계열: 인코딩 자동 감지 (UTF-8 → CP949)
    for enc in ("utf-8-sig", "utf-8", "cp949"):
        try:
            with open(path, "r", encoding=enc) as f:
                return f.read().strip()
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise UnicodeDecodeError("b", b"", 0, 1, "지원하지 않는 인코딩입니다 (UTF-8/CP949만 지원)")
